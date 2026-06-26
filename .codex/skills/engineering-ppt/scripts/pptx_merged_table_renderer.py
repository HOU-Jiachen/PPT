#!/usr/bin/env python3
"""Render DOCX table models as editable PPTX tables with merged cells."""

from __future__ import annotations

from typing import Any

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


def _rgb(hex_value: str):
    from pptx.dml.color import RGBColor

    value = hex_value.strip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _fit_text(text: str, limit: int) -> str:
    value = " ".join(str(text or "").split())
    if len(value) <= limit:
        return value
    boundary = max(value[:limit].rfind(mark) for mark in ("。", "；", "，", "、", " "))
    if boundary >= limit * 0.55:
        value = value[:boundary]
    else:
        value = value[:limit]
    return value.rstrip("，,；;、 ") + "。"


def _column_widths(model: dict[str, Any], total_width: float, col_count: int) -> list[float]:
    raw = [int(item) for item in model.get("grid_widths", []) if int(item or 0) > 0]
    if len(raw) != col_count:
        return [total_width / max(1, col_count)] * col_count
    total = sum(raw) or 1
    return [total_width * value / total for value in raw]


def render_merged_table(
    slide,
    model: dict[str, Any],
    x: float,
    y: float,
    w: float,
    h: float,
    colors: dict[str, str],
    max_rows: int | None = None,
    max_cols: int | None = None,
    table_min_pt: float = 12.0,
):
    """Render a full or lightly clipped table model, applying PPTX cell merges."""

    rows = model.get("rows") or []
    if max_rows:
        rows = rows[:max_rows]
    col_count = int(model.get("column_count") or (max(len(row) for row in rows) if rows else 1))
    if max_cols:
        col_count = min(col_count, max_cols)
    row_count = len(rows)
    if row_count == 0 or col_count == 0:
        return None

    shape = slide.shapes.add_table(row_count, col_count, Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table

    for row in table.rows:
        row.height = Inches(h / max(1, row_count))
    for column, width in zip(table.columns, _column_widths(model, w, col_count)):
        column.width = Inches(width)

    cell_limit = max(14, min(68, int((w / max(1, col_count)) * (h / max(1, row_count)) * 24)))
    origin_lookup = {(cell["row"], cell["col"]): cell for cell in model.get("cells", [])}

    for r_idx in range(row_count):
        row_values = rows[r_idx] if r_idx < len(rows) else []
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            text = row_values[c_idx] if c_idx < len(row_values) else ""
            origin = origin_lookup.get((r_idx, c_idx))
            if origin is not None:
                text = origin.get("text", text)
            cell.text = _fit_text(text, cell_limit)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.035)
            cell.margin_right = Inches(0.035)
            cell.margin_top = Inches(0.018)
            cell.margin_bottom = Inches(0.018)
            is_header = r_idx == 0 or (origin or {}).get("is_header")
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(colors["secondary"] if is_header else "FFFFFF")
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.font.name = "Microsoft YaHei"
                paragraph.font.size = Pt(max(table_min_pt, 10.5 if row_count > 8 else table_min_pt))
                paragraph.font.bold = bool(is_header)
                paragraph.font.color.rgb = _rgb(colors["ink"])

    for span in model.get("spans", []):
        r0 = int(span["row"])
        c0 = int(span["col"])
        rs = int(span.get("row_span", 1))
        cs = int(span.get("col_span", 1))
        r1 = r0 + rs - 1
        c1 = c0 + cs - 1
        if r0 < row_count and c0 < col_count and r1 < row_count and c1 < col_count and (rs > 1 or cs > 1):
            try:
                table.cell(r0, c0).merge(table.cell(r1, c1))
            except ValueError:
                pass

    return shape


def model_fits_native_table(model: dict[str, Any], max_rows: int = 10, max_cols: int = 7) -> bool:
    return int(model.get("row_count", 0)) <= max_rows and int(model.get("column_count", 0)) <= max_cols
